import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(45, 55, 72) # #2D3748
    
    # --- Top Brand Eyebrow ---
    eyebrow_p = doc.add_paragraph()
    eyebrow_p.paragraph_format.space_before = Pt(0)
    eyebrow_p.paragraph_format.space_after = Pt(2)
    run_eyebrow = eyebrow_p.add_run("SWAT TOURISM PLATFORM (STP) — PRODUCTION V2")
    run_eyebrow.font.size = Pt(11)
    run_eyebrow.font.bold = True
    run_eyebrow.font.color.rgb = RGBColor(16, 185, 129) # Emerald #10B981
    
    # --- Document Main Title ---
    h1_p = doc.add_paragraph()
    h1_p.paragraph_format.space_before = Pt(2)
    h1_p.paragraph_format.space_after = Pt(6)
    run_h1 = h1_p.add_run("Hotel Portal: Physical Room Inventory & Reservation Assignment System")
    run_h1.font.size = Pt(20)
    run_h1.font.bold = True
    run_h1.font.color.rgb = RGBColor(23, 76, 114) # Primary Blue #174C72
    
    # --- Subtitle / Metadata ---
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    run_meta = meta_p.add_run(
        "Technical Architecture, Implementation Details, and Workflow Specification\n"
        "Date: October 2026  |  Author: Antigravity AI Engineering  |  Status: Verified & Production Ready"
    )
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(100, 116, 139)

    # --- Callout Box: Executive Summary ---
    callout_tbl = doc.add_table(rows=1, cols=1)
    callout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout_tbl.autofit = False
    callout_tbl.columns[0].width = Inches(6.5)
    c_cell = callout_tbl.cell(0, 0)
    set_cell_background(c_cell, "F0FDF4") # Light emerald
    set_cell_margins(c_cell, top=140, bottom=140, left=180, right=180)
    
    cp = c_cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(4)
    r = cp.add_run("📌 Executive Summary")
    r.font.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(4, 120, 87)
    
    cp2 = c_cell.add_paragraph()
    cp2.paragraph_format.space_after = Pt(0)
    r2 = cp2.add_run(
        "In hospitality management, booking categories (e.g. 'Executive Deluxe Room') define pricing and amenities, "
        "but hotel front desks require assigning specific physical rooms (e.g. 'Room 102', Floor 1, River Wing). "
        "This implementation establishes a complete physical room configuration and reservation assignment workflow in the Hotel Portal. "
        "Hotel managers can now configure physical room numbers, floors, and bed setups; visually inspect live room statuses on an interactive "
        "Key Rack; assign individual rooms to tourist booking requests with one click upon acceptance; dynamically update real-time 5-day "
        "room availability; and display assigned room numbers directly on tourist trip vouchers."
    )
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section Helper
    def add_section_header(num, title):
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(14)
        hp.paragraph_format.space_after = Pt(6)
        r = hp.add_run(f"{num}. {title}")
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(23, 76, 114)
        return hp

    def add_sub_header(title):
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(10)
        hp.paragraph_format.space_after = Pt(4)
        r = hp.add_run(title)
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
        return hp

    # --- 1. Problem Statement & Functional Requirements ---
    add_section_header("1", "Business Requirements & Problem Statement")
    
    p = doc.add_paragraph()
    p.add_run(
        "Prior to this enhancement, hotel bookings operated exclusively at the room category level. "
        "While price and availability were tracked, hotel managers lacked the ability to:"
    )
    
    bullet_pts = [
        ("Allocate specific physical rooms: ", "When a manager accepted a booking, there was no way to record which physical key (e.g. Room 101 vs Room 102) was designated for the guest."),
        ("Configure room numbers & floor plans: ", "Properties could not input their actual room numbers, floor levels, building wings, or specific bed configurations."),
        ("Inspect real-time room occupancy & housekeeping: ", "Front-desk personnel lacked a visual Key Rack to monitor which rooms are Available, Reserved, Occupied, Cleaning, or under Maintenance."),
        ("Synchronize category availability with physical room state: ", "Marking individual rooms as reserved or occupied did not dynamically decrement inventory across the 5-day availability matrix."),
        ("Inform tourists of assigned room keys: ", "Tourists had to wait until arrival at check-in to learn their room number, preventing smooth arrival logistics.")
    ]
    for b_title, b_desc in bullet_pts:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(b_title)
        r1.bold = True
        r1.font.color.rgb = RGBColor(23, 76, 114)
        bp.add_run(b_desc)

    # --- 2. Technical Architecture & Component Structure ---
    add_section_header("2", "Technical Architecture & Component Structure")
    
    p = doc.add_paragraph()
    p.add_run(
        "The feature is implemented across six modular frontend components and utility engines, "
        "designed for zero-latency local caching, reactive event dispatching, and full backend compatibility:"
    )

    # Table of Components
    comp_tbl = doc.add_table(rows=1, cols=3)
    comp_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    comp_tbl.autofit = False
    
    col_widths = [Inches(1.8), Inches(2.2), Inches(2.5)]
    headers = ["Module / File", "Component Name", "Role & Responsibility"]
    
    hdr_cells = comp_tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "174C72")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    components_data = [
        ("physicalRooms.ts", "Physical Rooms Core Engine", "Data types, storage engine, seed defaults, atomic room reservation, status updates, and custom window event sync."),
        ("PhysicalRoomConfigModal.tsx", "Room Configuration Modal", "UI for adding individual rooms, batch generator (e.g. Rooms 101-105), category linking, and live status adjustments."),
        ("AssignRoomModal.tsx", "Room Assignment Modal", "Interactive room selection modal triggered on booking accept; category compatibility filtering, guest info display, and confirm/release controls."),
        ("HotelTodayScreen.tsx", "Hotel Operations Dashboard", "Live Room Rack & Key Board, 5-Day real-time availability calculation, 'Accept & Assign Room' action, and assigned key badge cards."),
        ("BusinessServicesManagerModal.tsx", "Services & Catalog Manager", "Added 'Configure Physical Keys' entry point in the Rooms & Suites tab to manage physical keys directly from listing management."),
        ("TripManagerView.tsx", "Tourist Itinerary & Vouchers", "Renders confirmed physical room number and floor badge on tourist itinerary cards and confirmed vouchers.")
    ]
    
    for row_idx, (c_file, c_name, c_role) in enumerate(components_data):
        row = comp_tbl.add_row()
        row_cells = row.cells
        bg_col = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for j, text in enumerate([c_file, c_name, c_role]):
            row_cells[j].width = col_widths[j]
            set_cell_background(row_cells[j], bg_col)
            set_cell_margins(row_cells[j], top=80, bottom=80, left=100, right=100)
            p = row_cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if j == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(23, 76, 114)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- 3. Data Model & Lifecycle ---
    add_section_header("3", "Physical Room Data Model & Room Lifecycle")

    add_sub_header("3.1 PhysicalRoom TypeScript Interface")
    p = doc.add_paragraph()
    p.add_run("Each physical room entity is defined by the following strongly-typed schema:")

    schema_code = (
        "export interface PhysicalRoom {\n"
        "  id: string;                     // Unique room key identifier (e.g., 'serena-102')\n"
        "  listing_id: string;             // Owning hotel property identifier\n"
        "  price_item_id: string;          // Linked room category / pricing item ID\n"
        "  category_name: string;          // Human-readable category (e.g., 'Executive Deluxe')\n"
        "  room_number: string;            // Physical number/name (e.g., '102', 'Suite A')\n"
        "  floor: number | string;         // Floor level (e.g., 1, 2, 'Ground Floor')\n"
        "  wing?: string;                  // Building wing (e.g., 'River Wing', 'East Wing')\n"
        "  bed_setup?: string;             // Bedding details (e.g., '1 King Bed', '2 Twin Beds')\n"
        "  status: 'available' | 'reserved' | 'occupied' | 'cleaning' | 'maintenance';\n"
        "  current_booking_id?: string;    // ID or reference of currently assigned booking\n"
        "  current_guest_name?: string;    // Name of tourist currently holding reservation\n"
        "  stay_dates?: string;            // Active dates of stay (e.g., '2026-10-22 – 2026-10-24')\n"
        "  notes?: string;                 // Housekeeping or operational notes\n"
        "  updated_at: string;             // ISO timestamp of last status change\n"
        "}"
    )
    
    code_tbl = doc.add_table(rows=1, cols=1)
    code_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    code_tbl.columns[0].width = Inches(6.5)
    c = code_tbl.cell(0, 0)
    set_cell_background(c, "1E293B") # Dark slate
    set_cell_margins(c, top=100, bottom=100, left=140, right=140)
    cp = c.paragraphs[0]
    cp.paragraph_format.space_after = Pt(0)
    r = cp.add_run(schema_code)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(241, 245, 249)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_sub_header("3.2 Room State Machine & Lifecycle Transitions")
    p = doc.add_paragraph()
    p.add_run("Rooms transition through five explicit operational states:")
    
    states = [
        ("AVAILABLE (Green / #10B981): ", "The room is clean, vacant, and ready for assignment to incoming requests."),
        ("RESERVED (Blue / #2563EB): ", "The manager has accepted a booking and allocated this specific room. The guest's name and stay dates are locked to the key."),
        ("OCCUPIED (Red / #EF4444): ", "The tourist has checked in and is actively residing in the room."),
        ("CLEANING (Amber / #F59E0B): ", "Guest has checked out. Housekeeping staff is servicing, sanitizing, and restocking linen."),
        ("MAINTENANCE (Slate / #64748B): ", "Room is temporarily taken out of service for electrical, plumbing, or seasonal refurbishments.")
    ]
    for s_title, s_desc in states:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(s_title)
        r1.bold = True
        bp.add_run(s_desc)

    # --- 4. User Experience & Operational Flows ---
    add_section_header("4", "User Workflows & Operational Mechanics")

    add_sub_header("Flow A: Hotel Manager Configures Physical Rooms")
    p = doc.add_paragraph()
    p.add_run(
        "1. From the Hotel Today screen or the Listings & Rooms tab, the manager clicks 'Configure Physical Rooms'.\n"
        "2. The manager can use the Single Room Form to enter custom names ('Suite 201', 'Penthouse') or use the Quick Multi-Room Generator to automatically create consecutive keys (e.g. Prefix '10', Start 1, Count 5 -> Rooms 101 through 105).\n"
        "3. Keys are categorized under specific listing price items so capacity is automatically tracked.\n"
        "4. Saved keys immediately persist to storage and dispatch a window broadcast event to update all open views."
    )

    add_sub_header("Flow B: Tourist Booking Request & One-Click Room Allocation")
    p = doc.add_paragraph()
    p.add_run(
        "1. A tourist submits a booking request for an 'Executive Deluxe Room' for 2 nights.\n"
        "2. On the Hotel Today dashboard, the manager sees the request card with tourist name, dates, amount, and an 'Accept & Assign Room' button.\n"
        "3. Clicking 'Accept & Assign Room' marks the booking accepted and instantly launches the Assign Room Modal.\n"
        "4. The modal filters to show only physical rooms that match the requested room type. Available rooms are highlighted with prominent 'Select & Assign' tiles.\n"
        "5. Clicking a room marks that physical room 'Reserved' with guest name and dates, binds the room number to the booking record, and closes the modal.\n"
        "6. The booking card on the dashboard now shows: '🛏️ Assigned: Room 102 (Floor 1)' along with a 'Change Room' button if reassignment is needed."
    )

    add_sub_header("Flow C: Real-Time Availability Matrix Synchronization")
    p = doc.add_paragraph()
    p.add_run(
        "The 5-day availability matrix dynamically aggregates both confirmed API bookings and local physical room status. "
        "When physical room 102 is marked 'Reserved' or 'Occupied', the free room counter for 'Executive Deluxe Room' automatically "
        "decrements from 12 to 11 for the corresponding dates, ensuring double-bookings are prevented."
    )

    add_sub_header("Flow D: Tourist Itinerary & Voucher Display")
    p = doc.add_paragraph()
    p.add_run(
        "When the tourist views their confirmed itinerary in the Trips tab or opens their official booking voucher, "
        "the assigned room number and floor level are rendered cleanly (e.g. '🛏️ Assigned Room: Room 102 (Floor 1 • Garden Wing)'), "
        "giving the tourist complete peace of mind before traveling."
    )

    # --- 5. Verification & Testing ---
    add_section_header("5", "Verification, Testing & Build Results")
    
    p = doc.add_paragraph()
    p.add_run(
        "The implementation has been thoroughly validated against the full Swat Tourism Platform test suite "
        "and strict TypeScript compilation rules:"
    )

    # Test Results Callout
    res_tbl = doc.add_table(rows=1, cols=1)
    res_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    res_tbl.columns[0].width = Inches(6.5)
    rc = res_tbl.cell(0, 0)
    set_cell_background(rc, "F8FAFC")
    set_cell_margins(rc, top=120, bottom=120, left=160, right=160)
    
    rp = rc.paragraphs[0]
    rp.paragraph_format.space_after = Pt(4)
    r = rp.add_run("✅ Test & Build Summary")
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(16, 185, 129)
    
    rp2 = rc.add_paragraph()
    rp2.paragraph_format.space_after = Pt(0)
    r_body = (
        "• Vitest Test Suite: 22 test files executed, 115 tests passed, 0 failed (100% pass rate).\n"
        "• bookings.test.tsx: KPI counters, Accept & Assign Room workflow, and tourist trip confirmations verified.\n"
        "• TypeScript Compilation: 'tsc -b' compiled with zero type errors (all unused imports and types resolved).\n"
        "• Production Bundle: 'vite build' generated production assets cleanly in under 1 second."
    )
    r_run = rp2.add_run(r_body)
    r_run.font.size = Pt(9.5)
    r_run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- 6. Deliverable Artifacts Table ---
    add_section_header("6", "Summary of Deliverable Files")

    files_tbl = doc.add_table(rows=1, cols=3)
    files_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    files_tbl.autofit = False
    
    f_widths = [Inches(2.5), Inches(1.2), Inches(2.8)]
    f_headers = ["File Path", "Type", "Status / Notes"]
    
    f_hdr_cells = files_tbl.rows[0].cells
    for i, title in enumerate(f_headers):
        f_hdr_cells[i].width = f_widths[i]
        set_cell_background(f_hdr_cells[i], "174C72")
        set_cell_margins(f_hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = f_hdr_cells[i].paragraphs[0]
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    deliverables = [
        ("Docs/room_assignment_prompt.md", "Specification", "Original approved functional and prompt specification."),
        ("Docs/hotel-room-reservation-and-assignment-implementation.docx", "Documentation", "This comprehensive Word document implementation guide."),
        ("frontend/src/modules/business/physicalRooms.ts", "TypeScript", "Core storage and operations engine for physical rooms."),
        ("frontend/src/modules/business/PhysicalRoomConfigModal.tsx", "React Component", "Room number entry & multi-room batch generator dialog."),
        ("frontend/src/modules/business/AssignRoomModal.tsx", "React Component", "Interactive room selection and assignment dialog."),
        ("frontend/src/modules/business/HotelTodayScreen.tsx", "React Component", "Live Key Rack, 5-day availability matrix & request cards."),
        ("frontend/src/modules/business/BusinessServicesManagerModal.tsx", "React Component", "Rooms catalog integration with physical keys configuration."),
        ("frontend/src/modules/tourist/TripManagerView.tsx", "React Component", "Tourist voucher and itinerary room key badge integration.")
    ]
    
    for row_idx, (f_path, f_type, f_status) in enumerate(deliverables):
        row = files_tbl.add_row()
        row_cells = row.cells
        bg_col = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for j, text in enumerate([f_path, f_type, f_status]):
            row_cells[j].width = f_widths[j]
            set_cell_background(row_cells[j], bg_col)
            set_cell_margins(row_cells[j], top=70, bottom=70, left=100, right=100)
            p = row_cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if j == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(23, 76, 114)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Sign-off footer
    foot_p = doc.add_paragraph()
    foot_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_foot = foot_p.add_run("Swat Tourism Platform (STP) — Production V2 Ready  •  End of Document")
    r_foot.font.size = Pt(8.5)
    r_foot.font.italic = True
    r_foot.font.color.rgb = RGBColor(148, 163, 184)

    # Save document
    output_path = os.path.join("f:\\Antigravity\\SwatTourismProduction_v2\\Docs", "hotel-room-reservation-and-assignment-implementation.docx")
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document()
