import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(45, 55, 72) # #2D3748
    
    # --- Title Banner ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Swat Tourism Platform — Production V2")
    run_title.font.size = Pt(13)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 185, 129) # Emerald #10B981
    
    h1_p = doc.add_paragraph()
    h1_p.paragraph_format.space_before = Pt(4)
    h1_p.paragraph_format.space_after = Pt(8)
    run_h1 = h1_p.add_run("AI-Assisted Tourist Journey: Comprehensive Architecture & Roadmap")
    run_h1.font.size = Pt(22)
    run_h1.font.bold = True
    run_h1.font.color.rgb = RGBColor(27, 79, 114) # Deep Blue #1B4F72
    
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(18)
    run_meta = meta_p.add_run("Core Design Principle: 'Less Input, More Output' for Mountain Travelers\n"
                              "Author: Google Deepmind Antigravity AI  |  Target: Tourist All-Functions  |  Status: Strategic Blueprint")
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(100, 116, 139)

    # --- Callout Box: Executive Vision ---
    callout_tbl = doc.add_table(rows=1, cols=1)
    callout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout_tbl.autofit = False
    callout_tbl.columns[0].width = Inches(6.5)
    c_cell = callout_tbl.cell(0, 0)
    set_cell_background(c_cell, "F0FDF4") # Light emerald
    set_cell_margins(c_cell, top=140, bottom=140, left=180, right=180)
    
    cp = c_cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(4)
    r = cp.add_run("🎯 Strategic Goal & Philosophy")
    r.font.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(4, 120, 87)
    
    cp2 = c_cell.add_paragraph()
    cp2.paragraph_format.space_after = Pt(0)
    r2 = cp2.add_run(
        "Mountain tourism in Swat presents unique frictions: unpredictable high-altitude weather, road pass blockages, "
        "intercity transit coordination, and complex family budgeting. Rather than forcing the tourist to fill lengthy forms, "
        "AI assistance transforms the platform into an intelligent co-pilot: providing maximum proactive value with minimal tourist input."
    )
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper for Section Titles
    def add_section_header(num, title):
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(16)
        hp.paragraph_format.space_after = Pt(6)
        r = hp.add_run(f"{num}. {title}")
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(27, 79, 114)
        
    def add_feature_card(title, tagline, details_list):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"• {title}: ")
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
        r2 = p.add_run(tagline)
        r2.font.italic = True
        r2.font.color.rgb = RGBColor(71, 85, 105)
        
        for item in details_list:
            bp = doc.add_paragraph()
            bp.paragraph_format.left_indent = Inches(0.25)
            bp.paragraph_format.space_before = Pt(1)
            bp.paragraph_format.space_after = Pt(2)
            run = bp.add_run(f"- {item}")
            run.font.size = Pt(10)

    # --- 1. Home Screen & Discovery ---
    add_section_header("1", "Home Screen & Discovery: AI 'Prompt-to-Trip' Co-Pilot")
    
    add_feature_card(
        "Natural Language Trip Generation (1-Prompt Itinerary)",
        "Zero-planning trip creation from natural language chat or speech.",
        [
            "Tourist types or speaks: '3-day family trip from Islamabad, under 60k PKR, want snow at Malam Jabba and riverside trout in Bahrain'.",
            "AI analyzes verified hotel room inventories, driver rates, and road clearance to assemble an instant day-by-day plan.",
            "Outputs direct 1-tap checkout with live Easypaisa / JazzCash payment breakdown."
        ]
    )
    add_feature_card(
        "Dynamic 'Vibe & Season' Recommender",
        "Instant matching based on real-time Swat seasonal pulses.",
        [
            "Automatically matches tourist mood chips (❄️ Snow & Ski, 🌊 Riverside Stays, 🏔️ 4x4 Alpine, 🏰 Heritage).",
            "Factors in active season: e.g. chairlift operation in Malam Jabba during snow season, cherry/apple blossom in Miandam, or trout season in Madyan.",
            "Reduces search friction by presenting curated spots upfront."
        ]
    )
    add_feature_card(
        "AI Mountain Packing & Preparation Checklist",
        "Contextual gear advisory tailored to chosen valleys.",
        [
            "Generates customized packing list based on altitude and forecast (e.g. thermal layers and waterproof boots for Mahodand Lake, motion sickness meds for Kalam pass).",
            "Includes cash ATM alerts (warning tourists that card terminals are scarce beyond Bahrain)."
        ]
    )

    # --- 2. Explore & Verified Listings ---
    add_section_header("2", "Explore & Verified Listings: AI Matcher & Menu Explainer")
    
    add_feature_card(
        "'Ask This Listing' AI Mini-Assistant",
        "Instant answers to terrain, heating, and vehicle accessibility questions.",
        [
            "Summarizes long amenity lists into 3 critical traveler facts (e.g. '24/7 Room Heating, Solar Backup, Private Riverfront').",
            "Answers road clearance queries: 'Can a sedan reach this hotel or is a 4x4 required?' based on elevation and terrain metadata."
        ]
    )
    add_feature_card(
        "Swati Cuisine & Traditional Menu Explainer",
        "Cultural culinary translation for authentic local dining.",
        [
            "Explains traditional dishes (Swati River Trout Fry, Shinwari Mutton Karahi, Kandahari Sesame Naan, Kalam Walnut Halwa).",
            "Provides spice ratings, family portion recommendations, and dietary filters (Halal, gluten-free, kid-friendly)."
        ]
    )
    add_feature_card(
        "Multilingual Voice Search (Urdu, English, Pashto)",
        "Hands-free voice discovery for travelers on the move.",
        [
            "Enables spoken search: 'Kalam mein 10,000 ke andar best family hotel dikhao' with automatic phonetic transcription.",
            "Applies filters instantly without typing on touchscreens during travel."
        ]
    )

    # --- 3. Trip Manager & Itinerary Planning ---
    add_section_header("3", "Trip Manager & Itinerary: AI Smart Route & Sequence Optimizer")
    
    add_feature_card(
        "Geographic Sequence Optimizer (Zero Backtracking)",
        "Intelligent stop reordering along the Swat River valley corridor.",
        [
            "Prevents inefficient travel loops (e.g. warns tourist against scheduling Mingora ➔ Kalam ➔ White Palace ➔ Kalam).",
            "Reorders itinerary stops geographically (Mingora ➔ Madyan ➔ Bahrain ➔ Kalam ➔ Mahodand) saving fuel and travel hours."
        ]
    )
    add_feature_card(
        "Realistic Mountain Drive Time Predictor",
        "Calculates winding mountain pass travel durations.",
        [
            "Replaces flat highway estimates with realistic mountain speeds factoring in elevation, hairpins, and viewpoint stops.",
            "Accurately estimates 3.5 hours for Mingora to Kalam and 2.5 hours for Kalam to Mahodand Lake 4x4 trail."
        ]
    )
    add_feature_card(
        "AI Budget Optimization & Smart Savings Coach",
        "Live financial recommendations across rooms, transport, and dining.",
        [
            "Analyzes estimated trip total and suggests actionable savings: 'Switching Night 2 to Executive Twin saves PKR 5,500 with zero comfort loss'.",
            "Suggests shared 4x4 group departures for popular lake trails."
        ]
    )
    add_feature_card(
        "Continuity & Booking Conflict Guard",
        "Automated audit of check-in, check-out, and transit windows.",
        [
            "Detects conflicting schedules: e.g. transport departing at 09:00 AM while hotel check-out is 12:00 PM.",
            "Flags missing meal windows or excessive travel segments without rest stops."
        ]
    )

    # --- 4. Safety, Road Status & Emergency SOS ---
    add_section_header("4", "Safety & Emergencies: Proactive Mountain Advisory AI")
    
    add_feature_card(
        "Predictive Snow & Weather Rerouting",
        "Proactive alerts before tourists reach blocked mountain passes.",
        [
            "Monitors weather forecasts and authority road condition updates.",
            "Alert: 'Heavy snowfall predicted at Malam Jabba Saturday 2:00 PM. AI recommends taking chairlift before noon or installing 4x4 snow chains.'",
            "1-tap alternative rerouting to open scenic valleys (e.g. Marghuzar White Palace or Bahrain Riverbank) if high passes close."
        ]
    )
    add_feature_card(
        "AI Emergency SOS Distress Dispatcher",
        "Automated bilingual rescue packet creation with GPS tagging.",
        [
            "When the tourist presses the SOS button, AI drafts a distress dispatch containing exact GPS coordinates, nearest landmark, and traveler medical notes.",
            "Prepares bilingual text in English and Urdu/Pashto for instant transmission to Rescue 1122, Swat Tourist Police, and emergency family contacts."
        ]
    )

    # --- 5. On-Trip WhatsApp & Voice Concierge ---
    add_section_header("5", "On-Trip Concierge: WhatsApp & Voice Travel Buddy")
    
    add_feature_card(
        "Direct WhatsApp Travel Assistant",
        "Lightweight, 2-second response travel buddy for low-bandwidth zones.",
        [
            "Tourists message their booking voucher or ask queries: 'Where is the nearest 24/7 pharmacy near Madyan?' or 'Send driver Sher Alam's phone number'.",
            "AI retrieves verified platform records and replies instantly with clickable telephone and map links."
        ]
    )
    add_feature_card(
        "Local Cultural Etiquette & Bargaining Guide",
        "Respectful guidance for engaging with local hosts and artisans.",
        [
            "Provides tips on photography etiquette in rural Swati villages, visiting historical Buddhist stupas, and fair pricing for hand-carved walnut crafts and Islampur shawls."
        ]
    )

    # --- 6. Strategic Implementation Roadmap ---
    add_section_header("6", "Strategic Phased Implementation Roadmap")
    
    tbl = doc.add_table(rows=5, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    col_widths = [Inches(1.2), Inches(2.3), Inches(1.5), Inches(1.5)]
    for row in tbl.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w
            
    headers = ["Phase", "AI Capability Module", "Implementation Effort", "Tourist Value Impact"]
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        set_cell_background(cell, "1B4F72") # Deep blue
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    data = [
        ("Phase A (Immediate)", "1-Prompt Natural Language Trip Builder\n• Parses budget, days, companions\n• Auto-populates multi-day itinerary", "Medium\n(LLM JSON Schema)", "⭐⭐⭐⭐⭐\n(Cuts planning time from 20 min to 30 sec)"),
        ("Phase B (Short-term)", "Smart Route & Sequence Optimizer\n• Eliminates travel backtracking\n• Realistic mountain drive times", "Low\n(Heuristic + Map Matrix)", "⭐⭐⭐⭐\n(Saves 3-5 hours on mountain roads)"),
        ("Phase C (Mid-term)", "Proactive Mountain Weather & SOS AI\n• Snow chains & pass alerts\n• Bilingual GPS emergency dispatch", "Medium\n(Weather API + GPS SOS)", "⭐⭐⭐⭐⭐\n(Critical life-safety feature)"),
        ("Phase D (Future)", "WhatsApp & Voice Travel Concierge\n• Voice search in Urdu/English\n• 24/7 low-bandwidth WhatsApp bot", "Medium-High\n(Twilio/WhatsApp API)", "⭐⭐⭐⭐\n(Peak on-road convenience)"),
    ]
    
    for row_idx, row_data in enumerate(data, start=1):
        bg = "F8FAFD" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = tbl.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(27, 79, 114)
            elif col_idx == 3 and "⭐⭐⭐⭐⭐" in text:
                r.font.bold = True
                r.font.color.rgb = RGBColor(4, 120, 87)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    
    # --- Conclusion & Technical Architecture ---
    add_section_header("7", "Technical Architecture & Data Privacy")
    conc_p = doc.add_paragraph()
    r = conc_p.add_run(
        "Architecture Guidelines:\n"
        "1. Hybrid Execution: Cloud LLM (Gemini 1.5 Flash / GPT-4o-mini) for rich natural language trip synthesis; "
        "client-side rule engine for instant route sequence optimization and offline emergency distress formatting.\n"
        "2. Strict Grounding: AI responses are grounded exclusively in verified business listings, published transport fleets, "
        "and official District Administration Swat road condition advisories to eliminate hallucinations.\n"
        "3. Zero Tourist Friction: All AI recommendations offer 1-tap confirmation pills ('+ Add to Day 1', 'Pay & Confirm Trip') "
        "adhering to the platform's high-speed, thumb-friendly interaction design."
    )
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(71, 85, 105)
    
    output_path = r"f:\Antigravity\SwatTourismProduction_v2\Docs\swat-tourism-tourist-ai-assistance-proposal.docx"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document()
