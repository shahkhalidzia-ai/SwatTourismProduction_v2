import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, text_prefix, text_body, border_color="0A4D68", fill_color="F0F7FA"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run1 = p.add_run(text_prefix + " ")
    run1.bold = True
    run1.font.name = "Calibri"
    run1.font.size = Pt(10)
    run1.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)
    
    run2 = p.add_run(text_body)
    run2.font.name = "Calibri"
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table_headers(table, headers, fill_hex="0A4D68", col_widths=None):
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], fill_hex)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if col_widths and i < len(col_widths):
            hdr_cells[i].width = Inches(col_widths[i])

def format_table_rows(table, data, col_widths=None, alternate_fill=True):
    for r_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        fill_hex = "F4F8FA" if (r_idx % 2 == 1 and alternate_fill) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], fill_hex)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            if col_widths and c_idx < len(col_widths):
                row_cells[c_idx].width = Inches(col_widths[c_idx])

def build_document():
    doc = Document()
    
    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "Swat Tourism Platform v2.0 | Architectural Gap Analysis & Advance Payment Policy"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.name = "Calibri"
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Confidential & Proprietary — Department of Tourism & Culture, Government of Khyber Pakhtunkhwa"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.runs[0].font.name = "Calibri"
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("Payment, Hotel Stay & Trip Completion")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("Architectural Gap Analysis, Advance Payment Rules & Production Roadmap")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x08, 0x83, 0x95)

    # Metadata Box
    meta_tbl = doc.add_table(rows=4, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False
    meta_data = [
        ("Platform Version:", "Swat Tourism Platform (STP) Production v2.0"),
        ("Target Domains:", "Advance Payment Models (Full/Partial), Hotel Front Desk, Trip Completion"),
        ("Prepared By:", "Lead Software Architect & AI Engineering Team"),
        ("Date & Status:", "October 2026 | Production Readiness Assessment (Updated)"),
    ]
    for r_idx, (k, v) in enumerate(meta_data):
        c0 = meta_tbl.cell(r_idx, 0)
        c1 = meta_tbl.cell(r_idx, 1)
        c0.width = Inches(1.8)
        c1.width = Inches(4.7)
        set_cell_background(c0, "EBF4F6")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.font.name = "Calibri"
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ---------------------------------------------------------
    # 1. Executive Summary
    # ---------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    run_h1 = h1.add_run("1. Executive Summary & Problem Statement")
    run_h1.font.name = "Calibri"
    run_h1.font.size = Pt(16)
    run_h1.font.bold = True
    run_h1.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)

    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.paragraph_format.space_after = Pt(6)
    r_exec = p_exec.add_run(
        "The Swat Tourism Platform (STP) has achieved rich discovery, booking request pipelines, transport fleet management, "
        "and physical room rack allocation. However, transitioning the platform into a commercial hospitality standard requires "
        "addressing key functional requirements: (1) multi-channel digital payments and flexible advance deposit options, "
        "(2) the physical hotel stay operational lifecycle (from front-desk check-in through in-stay folios and room turnover), "
        "and (3) automated trip completion triggers with verified multi-entity traveler feedback loops."
    )
    r_exec.font.name = "Calibri"
    r_exec.font.size = Pt(10)
    r_exec.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    add_callout(
        doc,
        "Core Architectural Verdict:",
        "The backend database schema already defines 'deposit_amount' and supports payment_status values ('unpaid', 'deposit_paid', 'paid', 'refunded'). "
        "However, the frontend UI currently lacks the interactive controls for tourists to select Full vs. Partial Deposit payment, and the "
        "Hotel Today Screen does not show front-desk staff the outstanding balance due upon guest arrival. Integrating this into the SwatPay modal "
        "and Hotel Check-In workflow is an essential deliverable."
    )

    # ---------------------------------------------------------
    # 2. Payment Architecture & Gap Analysis
    # ---------------------------------------------------------
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(4)
    run_h2 = h2.add_run("2. Payment Architecture & Gap Analysis (Tourist & Business)")
    run_h2.font.name = "Calibri"
    run_h2.font.size = Pt(16)
    run_h2.font.bold = True
    run_h2.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)

    p_pay = doc.add_paragraph()
    p_pay.paragraph_format.line_spacing = 1.15
    p_pay.paragraph_format.space_after = Pt(6)
    r_pay = p_pay.add_run(
        "Tourism dynamics across Swat Valley require a resilient, multi-rail payment architecture accommodating Easypaisa and JazzCash mobile "
        "wallets, SBP Raast instant QR codes, international credit/debit cards, and offline reservation guarantees for remote alpine areas."
    )
    r_pay.font.name = "Calibri"
    r_pay.font.size = Pt(10)
    r_pay.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    pay_tbl = doc.add_table(rows=1, cols=3)
    pay_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    pay_tbl.autofit = False
    pay_headers = ["Payment Sub-system", "Current Implementation State", "Identified Gaps & Required Upgrades"]
    pay_widths = [1.6, 2.4, 2.5]
    format_table_headers(pay_tbl, pay_headers, fill_hex="0A4D68", col_widths=pay_widths)

    pay_data = [
        (
            "Entire Trip Bundle Checkout\n(btn-pay-trip)",
            "Clicking 'Pay & Confirm Trip' in TripManagerView calls handleMockPayment in TouristShell, issuing a client toast and auto-updating trip status to 'confirmed'.",
            "CRITICAL GAP: Bypasses payment method selection entirely. No itemized checkout showing hotel rooms, rental cars, and tour guide totals. Needs unified 'SwatPay Modal'."
        ),
        (
            "Individual Booking Checkout\n(a-la-carte)",
            "EasypaisaPaymentModal exists for single bookings. Accepts mobile number and triggers backend API /pay/easypaisa with simulated/live merchant credential.",
            "HIGH GAP: Restricted exclusively to Easypaisa. Tourists cannot select JazzCash, Credit/Debit Cards, Raast, or Pay-on-Arrival from this modal."
        ),
        (
            "Backend Gateway Adapters\n(app/modules/bookings)",
            "PaymentProviderInterface with HMAC-SHA256 implemented for Easypaisa and JazzCash. Webhook listeners (/ipn-callback) verify incoming signatures.",
            "MEDIUM GAP: Lacks unified checkout endpoint (e.g. POST /bookings/checkout) accepting generic provider_name. Raast dynamic QR payload generator missing."
        ),
        (
            "Offline / Mountain Valley Continuity",
            "Bookings require digital verification or remain in unpaid/requested states. In cellular dead zones (Ushu/Mahodand), tourists cannot verify OTPs.",
            "HIGH GAP: Needs 'Pay on Arrival / Cash to Vendor' channel that generates guaranteed check-in vouchers with vendor contact instructions and offline PIN."
        ),
        (
            "Escrow & Merchant Settlement Ledger",
            "Payments flag bookings as 'paid' in database. Funds accumulate in platform account with no disaggregated accounting.",
            "MEDIUM GAP: No automated split between the 5% Tourism Development Levy and Net Vendor Payout (95%). Hotel dashboard lacks payout disbursement tracker."
        ),
    ]
    format_table_rows(pay_tbl, pay_data, col_widths=pay_widths)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---------------------------------------------------------
    # 3. ADVANCE PAYMENT POLICY & SPECIFICATION (NEW SECTION)
    # ---------------------------------------------------------
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(4)
    run_h3 = h3.add_run("3. Advance Payment Policy: Full vs. Partial Deposit vs. Pay-on-Arrival")
    run_h3.font.name = "Calibri"
    run_h3.font.size = Pt(16)
    run_h3.font.bold = True
    run_h3.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)

    p_adv = doc.add_paragraph()
    p_adv.paragraph_format.line_spacing = 1.15
    p_adv.paragraph_format.space_after = Pt(6)
    r_adv = p_adv.add_run(
        "A critical question in hospitality platforms is whether booking confirmation requires full upfront payment, a partial deposit, "
        "or zero advance payment. In the Swat Tourism Platform, the underlying database model already supports multiple commercial policies "
        "via 'deposit_amount' (Numeric) and 'payment_status' ('unpaid', 'deposit_paid', 'paid', 'refunded'). The platform establishes "
        "three distinct operational models:"
    )
    r_adv.font.name = "Calibri"
    r_adv.font.size = Pt(10)
    r_adv.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Table of Advance Payment Models
    adv_tbl = doc.add_table(rows=1, cols=4)
    adv_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    adv_tbl.autofit = False
    adv_headers = ["Payment Model", "Advance Required", "Booking Status & Confirmation Rules", "Front-Desk Balance Due"]
    adv_widths = [1.5, 1.4, 2.1, 1.5]
    format_table_headers(adv_tbl, adv_headers, fill_hex="0E8388", col_widths=adv_widths)

    adv_data = [
        (
            "1. Full Advance Payment\n(100% Upfront)",
            "100% of booking total charged at checkout via Easypaisa, JazzCash, Card, or Raast.",
            "• status: 'confirmed'\n• payment_status: 'paid'\n• Confirmed immediately upon gateway webhook confirmation.",
            "Balance Due: PKR 0\nGuest checked in without financial collection."
        ),
        (
            "2. Partial Advance Deposit\n(20% - 30% or 1st Night)",
            "Partial amount charged upfront (e.g. PKR 5,000 on a PKR 20,000 reservation).",
            "• status: 'confirmed'\n• payment_status: 'deposit_paid'\n• deposit_amount: recorded\n• Room is secured on hotel rack.",
            "Balance Due: PKR 15,000\nFront desk collects remaining balance at check-in or checkout."
        ),
        (
            "3. Pay on Arrival\n(0% Advance Guarantee)",
            "PKR 0 charged upfront. Reservation held via verified mobile / CNIC or card guarantee.",
            "• status: 'confirmed'\n• payment_status: 'unpaid'\n• Guaranteed voucher issued with check-in cancellation cut-off.",
            "Balance Due: PKR 20,000\nFull payment collected at hotel reception desk."
        ),
    ]
    format_table_rows(adv_tbl, adv_data, col_widths=adv_widths)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Detailed Business & Technical Rules for Advance Payments
    p_rules = doc.add_paragraph()
    p_rules.paragraph_format.line_spacing = 1.15
    p_rules.paragraph_format.space_after = Pt(4)
    r_rules_t = p_rules.add_run("Business Rules & Operational Implementation Requirements:\n")
    r_rules_t.bold = True
    r_rules_t.font.name = "Calibri"
    r_rules_t.font.size = Pt(11)
    r_rules_t.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)

    r_rules_b = p_rules.add_run(
        "1. Dynamic Tourist Selector in SwatPay Modal: The checkout modal must display clear payment choice options: "
        "'[X] Pay Full Amount (PKR 24,000 - Guaranteed Best Rate)' vs. '[ ] Pay 25% Deposit (PKR 6,000 now, PKR 18,000 at Hotel)' "
        "vs. '[ ] Pay on Arrival (PKR 0 now, PKR 24,000 at Venue)'.\n"
        "2. Hotel Front-Desk Balance Calculation: When a tourist arrives with payment_status = 'deposit_paid', the Hotel Today Screen "
        "must render a prominent badge: 'Deposit Paid: PKR 6,000 • Remaining Balance Due: PKR 18,000'. During check-in or check-out, "
        "the front desk records collection of this balance.\n"
        "3. High-Season Policy Enforcement: For high-demand luxury suites (e.g. Serena Heritage Suite or Malam Jabba Ski Resort in winter), "
        "hotel managers can toggle 'Advance Deposit Mandatory' in listing configuration, requiring at least a 25% non-refundable deposit.\n"
        "4. Cancellation & Refund Rules: If a tourist cancels a deposit-confirmed booking: cancellations made > 48 hours prior to check-in "
        "receive a 100% refund of deposit; cancellations < 48 hours forfeit the deposit to compensate the hotel for held room inventory."
    )
    r_rules_b.font.name = "Calibri"
    r_rules_b.font.size = Pt(9.5)
    r_rules_b.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---------------------------------------------------------
    # 4. Hotel Stay Lifecycle Gap Analysis
    # ---------------------------------------------------------
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(4)
    run_h4 = h4.add_run("4. Hotel Stay Operational Lifecycle & Room Operations")
    run_h4.font.name = "Calibri"
    run_h4.font.size = Pt(16)
    run_h4.font.bold = True
    run_h4.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)

    p_hotel = doc.add_paragraph()
    p_hotel.paragraph_format.line_spacing = 1.15
    p_hotel.paragraph_format.space_after = Pt(6)
    r_hotel = p_hotel.add_run(
        "A hotel reservation transitions through a physical on-property lifecycle: Arrival Verification -> Room Key Handover -> "
        "In-House Stay -> Incidentals/Folio Charges -> Check-Out & Billing -> Housekeeping Turnover. The current portal implements "
        "room allocation but lacks real-time operational execution."
    )
    r_hotel.font.name = "Calibri"
    r_hotel.font.size = Pt(10)
    r_hotel.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    hotel_tbl = doc.add_table(rows=1, cols=3)
    hotel_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hotel_tbl.autofit = False
    hotel_headers = ["Lifecycle Stage", "Current Implementation", "Operational Gap & Required Fix"]
    hotel_widths = [1.6, 2.3, 2.6]
    format_table_headers(hotel_tbl, hotel_headers, fill_hex="088395", col_widths=hotel_widths)

    hotel_data = [
        (
            "1. Guest Arrival &\nCheck-In Verification",
            "HotelTodayScreen renders 'Ready for check-in' badge under confirmed bookings. Physical room rack displays 'reserved' status.",
            "CRITICAL GAP: No interactive 'Check-In' button. Front desk staff cannot record guest arrival timestamp (actual_checkin_at), verify CNIC, or transition room to 'occupied'."
        ),
        (
            "2. Physical Key Assignment\n& Voucher Sync",
            "AssignRoomModal allows manager to choose room number (e.g. Room 204). Assignment is saved in browser localStorage.",
            "HIGH GAP: Assigned room is not stored in backend Booking model. Tourist digital voucher in TripManagerView does not show physical room number or WiFi access slip."
        ),
        (
            "3. In-House Stay &\nIncidental Folio",
            "No active in-house stay view. Booking total remains static based on room night rate.",
            "MEDIUM GAP: Cannot append room dining, laundry, or extra bed charges to guest stay. Staff have no guest folio summary for checkout settlement."
        ),
        (
            "4. Check-Out &\nBilling Settlement",
            "No check-out action button on booking card or room rack.",
            "CRITICAL GAP: Front desk cannot execute check-out. Booking does not transition to 'completed'. Outstanding folio balances cannot be collected."
        ),
        (
            "5. Housekeeping &\nRoom Turnover",
            "Room statuses ('available', 'reserved', 'occupied', 'cleaning') exist as a manual dropdown in the room rack settings.",
            "HIGH GAP: Check-out does not automatically flip room to 'cleaning'. Housekeeping staff receive no notification when a room is vacated and ready for sanitize/linen change."
        ),
    ]
    format_table_rows(hotel_tbl, hotel_data, col_widths=hotel_widths)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---------------------------------------------------------
    # 5. Trip Completion & Reviews Gap Analysis
    # ---------------------------------------------------------
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(4)
    run_h5 = h5.add_run("5. Trip Completion Lifecycle & Traveler Feedback Loop")
    run_h5.font.name = "Calibri"
    run_h5.font.size = Pt(16)
    run_h5.font.bold = True
    run_h5.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)

    p_trip = doc.add_paragraph()
    p_trip.paragraph_format.line_spacing = 1.15
    p_trip.paragraph_format.space_after = Pt(6)
    r_trip = p_trip.add_run(
        "Trip completion is the terminal state of the traveler journey that safeguards platform reputation and triggers vendor "
        "reputation updates. The current system relies on a manual user button click without automated date awareness or "
        "multi-vendor feedback aggregation."
    )
    r_trip.font.name = "Calibri"
    r_trip.font.size = Pt(10)
    r_trip.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    trip_tbl = doc.add_table(rows=1, cols=3)
    trip_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    trip_tbl.autofit = False
    trip_headers = ["Completion Component", "Current Implementation", "Deficiency & Needed Implementation"]
    trip_widths = [1.7, 2.3, 2.5]
    format_table_headers(trip_tbl, trip_headers, fill_hex="059669", col_widths=trip_widths)

    trip_data = [
        (
            "Trip State Automation\n(Temporal Engine)",
            "Trips remain in 'draft' or 'confirmed' status indefinitely until the tourist manually clicks 'Complete' in the header.",
            "HIGH GAP: System should auto-transition trip to 'in_progress' when start date is reached, and prompt for completion when end date passes and hotels are checked out."
        ),
        (
            "Multi-Entity Feedback\nLoop (Reviews)",
            "Clicking complete opens ReviewModal for a single booking item (usually the first hotel).",
            "CRITICAL GAP: A typical Swat trip bundles a Hotel, a 4x4 Jeep driver, and a local guide. Travelers need a multi-step review wizard evaluating each service independently."
        ),
        (
            "Trip Completion Dossier\n& Tax Invoices",
            "Individual booking vouchers are printable, but there is no unified journey summary.",
            "HIGH GAP: Missing downloadable 'Swat Journey Summary Dossier' compiling route map, all paid tax receipts, mileage logs, and digital tourism heritage certificate."
        ),
        (
            "Reputation & Rating Sync",
            "Reviews submit rating scores to local mock/API storage without recalculating verified vendor averages.",
            "MEDIUM GAP: Completed verified reviews should immediately update hotel rating_avg and rating_count badges on public exploration cards."
        ),
    ]
    format_table_rows(trip_tbl, trip_data, col_widths=trip_widths)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---------------------------------------------------------
    # 6. Target Solution Architecture
    # ---------------------------------------------------------
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(4)
    run_h6 = h6.add_run("6. Target Solution Architecture & Integration Plan")
    run_h6.font.name = "Calibri"
    run_h6.font.size = Pt(16)
    run_h6.font.bold = True
    run_h6.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)

    # Sub-item 1
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_after = Pt(3)
    r_s1_t = p_s1.add_run("A. Unified 'SwatPay' Checkout Modal (frontend/src/components/SwatPayModal.tsx):\n")
    r_s1_t.bold = True
    r_s1_t.font.name = "Calibri"
    r_s1_t.font.size = Pt(10.5)
    r_s1_t.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)
    r_s1_d = p_s1.add_run(
        "• Supports dual scope: Entire Trip Bundle (aggregated costs) and Single Booking.\n"
        "• Advance Payment Selector: Toggle between (a) Pay in Full (100%), (b) Pay Deposit (25%), or (c) Pay on Arrival.\n"
        "• Features 5 tabbed payment rails: Easypaisa, JazzCash, Visa/Mastercard 3DS, SBP Raast QR, and Pay-on-Arrival Guarantee.\n"
        "• Includes 1-Click 'Demo Auto-Fill' credentials for instant stakeholder demonstrations.\n"
        "• On confirmation, issues a digital QR-coded invoice and updates booking status to 'confirmed' with appropriate payment_status."
    )
    r_s1_d.font.name = "Calibri"
    r_s1_d.font.size = Pt(9.5)
    r_s1_d.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Sub-item 2
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_after = Pt(3)
    r_s2_t = p_s2.add_run("B. Interactive Hotel Front-Desk Controller (frontend/src/modules/business/HotelTodayScreen.tsx):\n")
    r_s2_t.bold = True
    r_s2_t.font.name = "Calibri"
    r_s2_t.font.size = Pt(10.5)
    r_s2_t.font.color.rgb = RGBColor(0x08, 0x83, 0x95)
    r_s2_d = p_s2.add_run(
        "• Balance Due Calculation: Renders 'Deposit Paid: PKR 6,000 | Balance Due: PKR 18,000' for partial payments.\n"
        "• Direct 'Check In Guest' button: Verifies reservation, stamps actual_checkin_at, shifts room from 'reserved' to 'occupied'.\n"
        "• Direct 'Check Out Guest' button: Produces final billing summary, shifts room to 'cleaning' status, and marks booking 'completed'.\n"
        "• Housekeeping Quick-Release: Staff click 'Mark Cleaned' on room card to return room to 'available' inventory.\n"
        "• Booking Data Sync: Updates backend Booking model with assigned room number so it renders on tourist vouchers."
    )
    r_s2_d.font.name = "Calibri"
    r_s2_d.font.size = Pt(9.5)
    r_s2_d.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Sub-item 3
    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.space_after = Pt(3)
    r_s3_t = p_s3.add_run("C. Trip Completion Engine & Multi-Service Feedback Wizard:\n")
    r_s3_t.bold = True
    r_s3_t.font.name = "Calibri"
    r_s3_t.font.size = Pt(10.5)
    r_s3_t.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
    r_s3_d = p_s3.add_run(
        "• Auto-triggers end-of-trip celebration modal when stay completes.\n"
        "• Sequenced multi-vendor review: Step 1 (Hotel Amenities & Cleanliness) -> Step 2 (Transport Driver Safety & Route) -> Step 3 (Activities/Guides).\n"
        "• Generates comprehensive 'Swat Journey Summary & Invoices' dossier with PDF export and print formatting."
    )
    r_s3_d.font.name = "Calibri"
    r_s3_d.font.size = Pt(9.5)
    r_s3_d.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ---------------------------------------------------------
    # 7. Technical Implementation Roadmap
    # ---------------------------------------------------------
    h7 = doc.add_paragraph()
    h7.paragraph_format.space_before = Pt(14)
    h7.paragraph_format.space_after = Pt(4)
    run_h7 = h7.add_run("7. Technical Implementation Roadmap & Verification Metrics")
    run_h7.font.name = "Calibri"
    run_h7.font.size = Pt(16)
    run_h7.font.bold = True
    run_h7.font.color.rgb = RGBColor(0x0A, 0x4D, 0x68)

    plan_tbl = doc.add_table(rows=1, cols=4)
    plan_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    plan_tbl.autofit = False
    plan_headers = ["Phase", "Milestone Deliverables", "Target Files", "Verification Metric"]
    plan_widths = [1.0, 2.5, 1.8, 1.2]
    format_table_headers(plan_tbl, plan_headers, fill_hex="0A4D68", col_widths=plan_widths)

    plan_data = [
        (
            "Phase 1",
            "Develop unified SwatPayModal.tsx with Full vs. Partial Deposit selector, 5 payment channels, and demo auto-fill.",
            "frontend/src/components/SwatPayModal.tsx",
            "100% responsive on Mobile (430px) & Desktop"
        ),
        (
            "Phase 2",
            "Wire SwatPayModal into TripManagerView.tsx (btn-pay-trip) and TouristShell.tsx.",
            "TripManagerView.tsx\nTouristShell.tsx",
            "Vitest suites pass without regression"
        ),
        (
            "Phase 3",
            "Implement Check-In, Check-Out, and Deposit Balance Due on HotelTodayScreen.",
            "HotelTodayScreen.tsx\nphysicalRooms.ts",
            "Balance due and room status reflect in real-time"
        ),
        (
            "Phase 4",
            "Sync assigned room number to backend Booking and render on tourist voucher.",
            "AssignRoomModal.tsx\nTripManagerView.tsx",
            "Voucher displays room number & QR code"
        ),
        (
            "Phase 5",
            "Implement multi-vendor review loop and consolidated Trip Summary Dossier.",
            "ReviewModal.tsx\nTripManagerView.tsx",
            "Verified ratings update listing averages"
        ),
    ]
    format_table_rows(plan_tbl, plan_data, col_widths=plan_widths)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Sign-off Block
    sign_p = doc.add_paragraph()
    sign_p.paragraph_format.line_spacing = 1.15
    sign_run = sign_p.add_run(
        "Approved for Implementation By: _____________________________        Date: _____________________\n"
        "Chief Technology Officer & Director General of Tourism, Swat Khyber Pakhtunkhwa"
    )
    sign_run.font.name = "Calibri"
    sign_run.font.size = Pt(9.5)
    sign_run.font.italic = True
    sign_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Save document
    output_dir = "Docs"
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "swat-tourism-payment-hotel-trip-completion-gap-analysis.docx")
    doc.save(output_path)
    print(f"Successfully generated updated Word Document at: {output_path}")

if __name__ == "__main__":
    build_document()
